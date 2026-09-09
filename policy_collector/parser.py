"""Extract authoritative metadata and full text; never invent missing fields."""
from __future__ import annotations
import io
import re
import urllib.parse
from datetime import date
from pathlib import Path
from bs4 import BeautifulSoup
from .models import Document

ATTACH_EXT = re.compile(r"\.(pdf|docx?|wps|xlsx?|zip|rar|ofd)$", re.I)

def date_value(text: str) -> str:
    m = re.search(r"((?:19|20)\d{2})\s*[-/年.]\s*(\d{1,2})\s*[-/月.]\s*(\d{1,2})", text or "")
    if m:
        try:
            return date(*map(int, m.groups())).isoformat()
        except ValueError:
            pass
    return ""


def attachment_links(soup, page_url: str) -> list[dict]:
    result, seen = [], set()
    _NAME_PARAMS = {"filename", "name", "file", "fileurl", "url", "attachname", "filename1", "downloadname"}
    for a in soup.find_all("a"):
        href = a.get('href','').strip()
        name = a.get('download') or a.get_text(' ',strip=True)
        # 重庆规范性文件的公开下载按钮：只解析两个字符串参数，不执行网页脚本。
        click = re.fullmatch(r"\s*downloadFj\(\s*'([^']+)'\s*,\s*'([^']+)'\s*\);?\s*", a.get('onclick',''))
        if click:
            name, href = click.groups()
        if not href: continue
        url = urllib.parse.urljoin(page_url, href)
        parsed = urllib.parse.urlsplit(url)
        if parsed.scheme not in ("http", "https", "file"):
            continue
        # 扩展名定位三级：① 路径本身（常规附件）；② 下载网关 query 的文件名参数
        # （如浙江 JPaaS download?fileUrl=…&fileName=通知无痕迹稿.wps）；③ 链接文本兜底。
        m = ATTACH_EXT.search(urllib.parse.unquote(parsed.path))
        if not m and parsed.query:
            for kv in parsed.query.split('&'):
                k, _, v = kv.partition('=')
                if k.lower() in _NAME_PARAMS:
                    decoded = urllib.parse.unquote(v)
                    if ATTACH_EXT.search(decoded):
                        m = ATTACH_EXT.search(decoded)
                        if not name and not href.lower().endswith(('.pdf','.doc','.docx','.wps')):
                            name = decoded  # 网关文件名即附件名
                        break
        if not m:
            m = ATTACH_EXT.search(urllib.parse.unquote(name))
        if not m or url in seen:
            continue
        seen.add(url)
        path = urllib.parse.unquote(parsed.path)
        result.append({"name": name or Path(path).name,
                       "url": url, "fmt": m.group(1).lower()})
    return result


class Parser:
    _WENHAO_RE = re.compile(r"([A-Za-z\u4e00-\u9fff]{2,24}[〔\[【（(]\s*\d{4}\s*[〕\]】）)]\s*\d{1,6}\s*号)")

    def parse(self, data: bytes, fmt: str, page_url: str = "", origin_path: str = "") -> Document:
        fmt = fmt.lower().lstrip(".")
        if data.startswith(b"%PDF-"):
            fmt = "pdf"
        if fmt in ("html", "htm", "shtml"):
            return self._html(data, page_url, origin_path)
        doc = Document(page_url=page_url, origin_path=origin_path)
        try:
            if fmt == "pdf":
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(data))
                if len(reader.pages) > 500:
                    doc.parse_error = "PDF超过500页，保留原件待处理"
                    return doc
                pages = [p.extract_text() or "" for p in reader.pages]
                doc.content = "\n\n".join(f"[第{i+1}页]\n{t}" for i,t in enumerate(pages) if t.strip())
                if any(not t.strip() for t in pages):
                    doc.parse_error = "PDF包含无文本页，需OCR复核"
            elif fmt == "docx":
                from docx import Document as WordDocument
                from docx.oxml.ns import qn
                d = WordDocument(io.BytesIO(data))
                blocks = []
                for block in d.element.body:
                    if block.tag in (qn("w:p"), qn("w:tbl")):
                        blocks.append(" ".join(x.text for x in block.iter(qn("w:t")) if x.text))
                doc.content = "\n".join(blocks).strip()
            elif fmt in ("txt",):
                doc.content = data.decode("utf-8-sig")
            else:
                doc.parse_error = f"暂不支持{fmt}正文解析，原件保留"
                return doc
        except Exception as e:
            doc.parse_error = f"{fmt}解析失败：{type(e).__name__}"
            return doc
        if not doc.content.strip():
            doc.parse_error = doc.parse_error or "正文为空，可能需要OCR"
        self._fill_meta(doc, doc.content)
        if not doc.title:
            lines = [x.strip() for x in doc.content.splitlines() if x.strip() and not x.startswith("[第")]
            doc.title = lines[0][:300] if lines else ""
        return doc

    def _html(self, data: bytes, page_url: str, origin_path: str) -> Document:
        soup = BeautifulSoup(data, "lxml")
        meta = {(m.get("name") or m.get("property") or "").lower(): m.get("content", "").strip()
                for m in soup.find_all("meta")}
        title = meta.get("articletitle", "") or meta.get("og:title", "")
        if not title:
            h = soup.find("h1") or soup.find("h2", class_=re.compile("title", re.I)) or soup.title
            title = h.get_text(" ", strip=True) if h else ""
        doc = Document(page_url=page_url, origin_path=origin_path, title=title[:500])
        if any(x in title for x in ("访问验证", "安全验证", "访问被拒绝", "Access Denied", "验证码")):
            doc.parse_error = "网页返回访问验证，需维护来源"
            return doc
        # Capture metadata before scripts and toolbars are removed.
        for key in ("pubdate", "publishdate", "article:published_time", "firstpublishedtime", "date", "发布日期"):
            doc.page_date = date_value(meta.get(key, ""))
            if doc.page_date:
                break
        doc.wenhao = meta.get("文号", "") or meta.get("docnumber", "")
        doc.issuing_authority = meta.get("发文机关", "") or meta.get("issuingauthority", "")
        doc.doc_date = date_value(meta.get("成文日期", "") or meta.get("documentdate", ""))
        all_text = soup.get_text(" ", strip=True)
        if not doc.page_date:
            m = re.search(r"(?:发布时间|发布日期|发布[日时]间|日期)\s*[:：]\s*([\d年月日/ .-]{8,24})", all_text)
            if m:
                doc.page_date = date_value(m.group(1))
        if not doc.doc_date:
            m = re.search(r"(?:成文日期|生成日期|成文时间)\s*[:：]?\s*([\d年月日/ .-]{8,24})", all_text)
            if m:
                doc.doc_date = doc.doc_date or date_value(m.group(1))
        if not doc.wenhao:
            # 文件编号/发文字号/文号：如浙江页面"文件编号：浙发改能源〔2026〕116号"
            m = re.search(r"(?:文件编号|发文字号|文号)\s*[:：]?\s*(\S{2,45}号)", all_text)
            if m:
                doc.wenhao = m.group(1)
        publisher = ""
        if not doc.issuing_authority:
            m = re.search(r"(?:发布机构|发文机关)\s*[:：]\s*(\S{2,30}?(?:发改委|政府|厅|局|委|办公室|部))", all_text)
            if m:
                publisher = m.group(1)
        # 文号可能在正文容器外，也可能被 Word 内联标签切碎。
        # 仅接受独立元数据/段落中的完整文号，避免把正文引用当成本文件文号。
        if not doc.wenhao:
            for tag in soup.select("p, span, td, .rules_tit1"):
                value = tag.get_text("", strip=True)
                if self._WENHAO_RE.fullmatch(value):
                    doc.wenhao = re.sub(r"\s+", "", value)
                    break
        doc.attachments = attachment_links(soup, page_url)
        for tag in soup.select("p"):
            if not tag.find(["p", "div", "table", "br", "a", "img"]):
                tag.string = tag.get_text("", strip=True)
        from .site_adapters import get_detail_adapter
        adapter = get_detail_adapter(page_url)
        parsed = adapter(soup, page_url) if adapter else None
        if parsed and parsed.content:
            doc.title = title or parsed.title
            # NDRC <title> contains site decoration; its adapter strips it.
            if not meta.get("articletitle") and not meta.get("og:title"):
                doc.title = parsed.title or doc.title
            doc.content = parsed.content
            doc.page_date = doc.page_date or parsed.page_date
        else:
            for tag in soup.select("script,style,noscript,header,footer,nav,.toolbar,.share,.print,.breadcrumb"):
                tag.decompose()
            main = None
            for selector in ("#UCAP-CONTENT", "#zoom", ".TRS_Editor", ".TRS_UEDITOR", ".tyxlContent", "#zoomcon", ".article-content",
                             ".tys-main-zt-show", ".article_con", ".art_con", ".bt_content", "#content", "article", ".content"):
                candidate = soup.select_one(selector)
                if candidate and len(candidate.get_text(strip=True)) > 20:
                    main = candidate
                    break
            if main is None:
                main = soup.body or soup
                doc.parse_error = '未定位正文容器，整页文本仅供复核，需维护来源适配'
            doc.content = re.sub(r"\n{3,}", "\n\n", main.get_text("\n", strip=True)).strip()
        self._fill_meta(doc, doc.content)
        doc.issuing_authority = doc.issuing_authority or publisher
        return doc

    def _fill_meta(self, doc: Document, text: str) -> None:
        # Prefer a standalone issuing number over a cited policy in a paragraph.
        if not doc.wenhao:
            lines = text.splitlines()
            own = next((self._WENHAO_RE.fullmatch(x.strip()) for x in lines[:30]
                        if self._WENHAO_RE.fullmatch(x.strip())), None)
            m = own or self._WENHAO_RE.search(doc.title)
            if m:
                doc.wenhao = re.sub(r"\s+", "", m.group(1))
        if not doc.doc_date or not doc.issuing_authority:
            # Only a standalone signature date; never the first effective/cited date.
            dates = list(re.finditer(r"(?m)^\s*((?:19|20)\d{2}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日)\s*$", text))
            if dates:
                m = dates[-1]
                doc.doc_date = doc.doc_date or date_value(m.group(1))
                if not doc.issuing_authority:
                    before = [x.strip() for x in text[:m.start()].splitlines() if x.strip()]
                    issuers = []
                    for line in reversed(before[-20:]):
                        if len(line) <= 45 and re.search(r"(?:人民政府|委员会|部|局|厅|办公室|发展改革委|中心)$", line):
                            issuers.insert(0, line)
                        else:
                            break
                    doc.issuing_authority = "、".join(issuers)
