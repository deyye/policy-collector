import io
import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from types import SimpleNamespace
import pytest
from policy_collector.agent import PolicyAgent
from policy_collector.config import AppConfig, SourceConfig, LLMConfig
from policy_collector.collector import FetchResult
from policy_collector.models import Document, Classification
from policy_collector.pipeline import Pipeline
from policy_collector.classifier import Classifier
from policy_collector.webapp import create_app


@pytest.mark.parametrize('choice,retries',[('retry_materials',1),('review_materials',0),('execute_shell',1)])
def test_agent_bounded_tool_choice(choice,retries):
    client=SimpleNamespace(available=True,usage={'input_tokens':5,'output_tokens':2},chat_json=lambda *a:{'action':choice})
    events=[];called=[]
    agent=PolicyAgent(SimpleNamespace(llm=SimpleNamespace(client=client)),lambda *a:events.append(a))
    doc=Document(attachments=[{'url':'https://agency.gov.cn/a.pdf','error':'ReadTimeout','parse_status':'failed'}])
    agent.prepare(doc,lambda urls:called.append(urls))
    assert len(called)==retries and agent.usage['input_tokens']==5
    assert all('execute_shell' not in str(e) for e in events)


def test_no_key_local_planner_and_no_retry_for_403():
    events=[];called=[]
    agent=PolicyAgent(SimpleNamespace(llm=SimpleNamespace(client=SimpleNamespace(available=False))),lambda *a:events.append(a))
    doc=Document(attachments=[{'url':'https://agency.gov.cn/a.pdf','error':'HTTP 403','parse_status':'failed'}])
    agent.prepare(doc,lambda urls:called.append(urls))
    assert not called and events[-1][1]=='attention'
    doc.attachments[0]['error']='ReadTimeout'
    agent.prepare(doc,lambda urls:called.append(urls))
    assert len(called)==1


def test_complete_document_has_no_extra_model_planning():
    client=SimpleNamespace(available=True,chat_json=lambda *a:pytest.fail('No planner call needed'))
    cls=SimpleNamespace(llm=SimpleNamespace(client=client,available=True),classify=lambda *a:Classification(is_investment_policy='no',need_review=True))
    agent=PolicyAgent(cls,lambda *a:None)
    agent.prepare(Document(content='完整正文'),lambda *a:pytest.fail('No repair needed'))
    assert agent.classify(Document(content='完整正文')).is_investment_policy=='pending'


def cfg_for(tmp_path):
    cfg=AppConfig.load();cfg.data_dir=tmp_path;cfg.db_path=tmp_path/'db';cfg.downloads_dir=tmp_path/'raw';cfg.llm.enabled=False
    cfg.sources={'test':SourceConfig(name='test',site='合成测试',list_url='https://agency.gov.cn/policy/',max_pages=1)}
    return cfg


def test_agent_repairs_transient_attachment_and_records_events(tmp_path,monkeypatch):
    from docx import Document as Word
    word=Word();word.add_paragraph('投资项目给予专项资金支持，明确资金使用和监督管理要求。');buf=io.BytesIO();word.save(buf)
    cfg=cfg_for(tmp_path);p=Pipeline(cfg);calls=[]
    def fetch(url):
        calls.append(url)
        return FetchResult(ok=False,error='ReadTimeout') if len(calls)==1 else FetchResult(content=buf.getvalue())
    monkeypatch.setattr(p.collector,'fetch',fetch)
    raw='<meta charset="utf-8"><h1>投资项目支持办法</h1><article>投资项目给予专项资金支持，明确资金使用和监督管理要求。<a href="a.docx">政策附件</a></article>'.encode()
    try:
        result=p.ingest_url(cfg.sources['test'],'https://agency.gov.cn/policy/1.html',raw=raw,prefer='rule')
        assert len(calls)==2 and result.attachments_failed==0 and result.ingested==1
        row=p.db.query_policies()[0];events=p.db.policy_agent_events(row['id'])
        assert any(e['action']=='repair' and e['status']=='done' for e in events)
        assert p.db.list_attachments(row['id'])[0]['parse_status']=='ok'
    finally:p.close()


def test_actual_local_compatible_api_without_key(tmp_path,monkeypatch):
    for key in ('LLM_API_KEY','DASHSCOPE_API_KEY'):monkeypatch.delenv(key,raising=False)
    received=[]
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            received.append((self.path,self.headers.get('Authorization'),json.loads(self.rfile.read(int(self.headers['Content-Length'])))))
            answer={'is_investment_policy':'yes','doc_type':'正式政策','category':['guarantee'],
                    'category_reason':'项目用地保障','evidence':'优先保障重大项目新增建设用地','category_evidence':{'guarantee':'优先保障重大项目新增建设用地'},'need_review':False,'confidence':.95}
            body=json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(answer,ensure_ascii=False)}}],'usage':{'prompt_tokens':20,'completion_tokens':15}}).encode()
            self.send_response(200);self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(body)
        def log_message(self,*args):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:
        cfg=cfg_for(tmp_path);cfg.llm=LLMConfig(enabled=True,base_url=f'http://127.0.0.1:{server.server_port}/v1',model='mock-model',api_key_env='UNSET_TEST_KEY',retries=0)
        result=PolicyAgent(Classifier(cfg),lambda *a:None).classify(Document(title='项目用地保障办法',content='优先保障重大项目新增建设用地。'))
        assert result.method=='llm' and not result.need_review and result.category=='guarantee'
        assert received[0][0]=='/v1/chat/completions' and received[0][1] is None
        assert result.input_tokens==20
    finally:server.shutdown();server.server_close();thread.join(timeout=2)


def test_run_progress_and_ui_routes(tmp_path):
    cfg=cfg_for(tmp_path);p=Pipeline(cfg)
    try:
        result=p.run_demo();r=p.db.list_runs()[0];progress=json.loads(r['progress'])
        assert progress['completed']==progress['total'] and progress['stage']=='finished'
        assert p.db.run_events(r['run_id'])
        client=create_app(cfg).test_client()
        for path in ['/', '/runs', '/runs/'+r['run_id'], '/sources', '/settings/model']:
            response=client.get(path);assert response.status_code==200
        body=client.get('/runs/'+r['run_id']).get_data(as_text=True)
        assert '处理时间线' in body and '本地规则' in body
    finally:p.close()
