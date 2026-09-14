"""Regression tests use synthetic originals, not business acceptance labels."""
import io
import json
import hashlib
import zipfile
from pathlib import Path
from types import SimpleNamespace
import pytest
from docx import Document as Word
from policy_collector.config import AppConfig,SourceConfig
from policy_collector.pipeline import Pipeline
from policy_collector.models import Document,Classification
from policy_collector.collector import FetchResult
from policy_collector.parser import Parser
from policy_collector.quality import repair_materials,attachment_report
from policy_collector.sampling import prepare_sample,score_review
from scripts.validate_batch import run_batch

@pytest.fixture
def pipe(tmp_path):
    cfg=AppConfig.load();cfg.db_path=tmp_path/'policy.db';cfg.data_dir=tmp_path;cfg.downloads_dir=tmp_path/'downloads'
    cfg.sources={'test':SourceConfig(name='test',site='合成官网',list_url='https://agency.gov.cn/policy/')}
    cfg.fetch.retries=0;cfg.fetch.request_interval_seconds=0
    p=Pipeline(cfg);yield p;p.close()

def word(text='项目投资给予资金支持，并实施项目全过程监管。'):
    b=io.BytesIO();d=Word();d.add_paragraph(text);d.save(b);return b.getvalue()

def html(attachment='a.docx',title='项目投资管理办法'):
    return f'<meta charset="utf-8"><h1>{title}</h1><article><p>项目投资安排与资金支持详见下文及附件。</p><a href="{attachment}">附件{attachment}</a></article>'.encode()

@pytest.mark.parametrize('human',[None,'confirm','adjust','reject'])
def test_same_bytes_reparse_updates_without_revision_and_preserves_human(pipe,monkeypatch,human):
    original=pipe.parser.parse
    def old_parse(data,fmt,**kw):
        if fmt=='docx':return Document(parse_error='旧解析器不支持')
        return original(data,fmt,**kw)
    monkeypatch.setattr(pipe.parser,'parse',old_parse)
    raw=word();monkeypatch.setattr(pipe.collector,'fetch',lambda u:FetchResult(content=raw,final_url=u))
    url='https://agency.gov.cn/policy/1.html';src=pipe.cfg.sources['test']
    pipe.ingest_url(src,url,raw=html(),prefer='rule');p=pipe.db.query_policies()[0]
    if human:pipe.db.audit(p['id'],human,['guarantee'] if human=='adjust' else [],'合成测试人工决定')
    before=pipe.db.get_policy(p['id'])
    monkeypatch.setattr(pipe.parser,'parse',original)
    result=repair_materials(pipe,local_only=True,prefer='rule')
    after=pipe.db.get_policy(p['id']);a=pipe.db.list_attachments(p['id'])[0]
    assert result['stats']['reparsed']==1 and len(pipe.db.versions(p['policy_key']))==1
    assert a['parse_status']=='ok' and '全过程监管' in a['parsed_text']
    if human:
        assert all(after[k]==before[k] for k in ('review_status','category','is_investment_policy','reason','evidence'))
        assert after['parse_requires_review']==1
        assert any(r['id']==p['id'] for r in pipe.db.query_policies(review_status='pending'))
    else:assert result['stats']['reclassified']==1
    rerun=repair_materials(pipe,local_only=True,prefer='rule',policy_id=p['id'])
    assert rerun['stats']['reparsed']==0 and rerun['stats']['reclassified']==0
    assert sum(h['action']=='reparse' for h in pipe.db.review_history(p['id']))==1


def test_repair_reclassifies_stale_material_verdict(pipe,monkeypatch):
    """结论已过期（材料已补齐但结论仍停在「待补材料」）时必须重判，且重复跑要幂等。

    这是 repair 存在的意义：材料变了结论要跟着变。曾经的缺陷是 repair 不传 reclassify，
    附件虽已解析出正文，条目却永远挂在待补材料队列里（实测 14 条全部落进 duplicates）。
    """
    raw=word();monkeypatch.setattr(pipe.collector,'fetch',lambda u:FetchResult(content=raw,final_url=u))
    src=pipe.cfg.sources['test'];url='https://agency.gov.cn/policy/1.html'
    pipe.ingest_url(src,url,raw=html(),prefer='rule');p=pipe.db.query_policies()[0]
    assert pipe.db.list_attachments(p['id'])[0]['parse_status']=='ok'
    # 人为制造「结论过期」：材料已解析完整，结论却仍标为待补材料
    with pipe.db.tx() as c:
        c.execute("UPDATE policies SET todo_type='material',need_review=0 WHERE id=?",(p['id'],))
    result=repair_materials(pipe,local_only=True,prefer='rule',policy_id=p['id'])
    assert result['stats']['reclassified']==1
    assert pipe.db.get_policy(p['id'])['todo_type']!='material'
    rerun=repair_materials(pipe,local_only=True,prefer='rule',policy_id=p['id'])
    assert rerun['stats']['reclassified']==0


def test_attachment_dead_link_does_not_freeze_verdict(pipe,monkeypatch):
    """站点已失效（404）的附件不该把结论永久冻在「待补材料」；暂时性失败仍应挡住。

    实测事故：某政策 2 个附件，1 个 HTTP 200（22106 字节）、1 个 HTTP 404（文件已从站点移除）。
    旧逻辑把"任一附件失败"一律当成材料缺口 → 该政策永远挂在待补材料队列，队列再也排不空，
    每次 repair 还白跑一次下载。区分永久/暂时失败后：有可用材料就照常判定并留痕。
    """
    raw=word()
    def fetch(url):
        if 'gone' in url:return FetchResult(ok=False,error='HTTP 404')
        if 'flaky' in url:return FetchResult(ok=False,error='HTTP 503')
        return FetchResult(content=raw,final_url=url)
    monkeypatch.setattr(pipe.collector,'fetch',fetch)
    src=pipe.cfg.sources['test']
    pipe.ingest_url(src,'https://agency.gov.cn/policy/1.html',raw=html().replace(
        b'</article>',b'<a href="gone.docx">gone.docx</a></article>'),prefer='rule')
    pipe.ingest_url(src,'https://agency.gov.cn/policy/2.html',raw=html().replace(
        b'</article>',b'<a href="flaky.docx">flaky.docx</a></article>'),prefer='rule')
    rows={r['page_url']:dict(r) for r in pipe.db._conn.execute('SELECT * FROM policies')}
    dead=rows['https://agency.gov.cn/policy/1.html']
    assert dead['todo_type']!='material','死链不该把结论冻在待补材料'
    assert '已失效' in (dead['reviewer_hint'] or '')
    flaky=rows['https://agency.gov.cn/policy/2.html']
    assert flaky['todo_type']=='material','暂时性失败仍应等补采'


def test_failed_sibling_does_not_discard_successful_reparse(pipe,monkeypatch):
    original=pipe.parser.parse;raw=word()
    def fetch(url):return FetchResult(ok=False,error='503') if 'missing' in url else FetchResult(content=raw,final_url=url)
    monkeypatch.setattr(pipe.collector,'fetch',fetch)
    monkeypatch.setattr(pipe.parser,'parse',lambda data,fmt,**kw:Document(parse_error='旧版本') if fmt=='docx' else original(data,fmt,**kw))
    page=html().replace(b'</article>',b'<a href="missing.pdf">missing.pdf</a></article>')
    pipe.ingest_url(pipe.cfg.sources['test'],'https://agency.gov.cn/policy/1.html',raw=page,prefer='rule')
    p=pipe.db.query_policies()[0];monkeypatch.setattr(pipe.parser,'parse',original)
    result=repair_materials(pipe,local_only=True,prefer='rule')
    assert result['stats']['reparsed']==1
    assert pipe.db.list_attachments(p['id'])[0]['parse_status']=='ok'
    report=attachment_report(pipe.db)
    assert report['download_ok']==1 and report['parse_complete']==1 and report['affected_policies']==1


def test_new_attachment_bytes_create_policy_revision(pipe,monkeypatch):
    payload=[word()];monkeypatch.setattr(pipe.collector,'fetch',lambda u:FetchResult(content=payload[0],final_url=u))
    src=pipe.cfg.sources['test'];url='https://agency.gov.cn/policy/1.html'
    pipe.ingest_url(src,url,raw=html(),prefer='rule');p=pipe.db.query_policies()[0]
    payload[0]=word('项目投资实施细则新版本，增加信贷保障。')
    assert pipe.ingest_url(src,url,raw=html(),prefer='rule').updated==1
    assert len(pipe.db.versions(p['policy_key']))==2


def test_ofd_manifest_order_and_graphics_gap():
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w') as z:
        z.writestr('OFD.xml','<OFD><DocBody><DocRoot>Doc/Document.xml</DocRoot></DocBody></OFD>')
        z.writestr('Doc/Document.xml','<Document><Pages><Page BaseLoc="b.xml"/><Page BaseLoc="a.xml"/></Pages></Document>')
        z.writestr('Doc/a.xml','<Page><TextCode>第二页内容</TextCode><ImageObject/></Page>')
        z.writestr('Doc/b.xml','<Page><TextCode>第一页政策正文</TextCode></Page>')
    d=Parser().parse(b.getvalue(),'ofd')
    assert d.content.index('第一页')<d.content.index('第二页')
    assert d.total_pages==2 and d.parsed_pages==1 and d.parse_error


def test_wps_with_docx_signature_is_readable():
    assert '资金支持' in Parser().parse(word(),'wps').content


def test_audit_excluded_reference_miss_and_frozen_material(pipe,tmp_path):
    pipe.sync_sources();sid=pipe.db.get_source('test')['id']
    pipe.db.add_fetch(sid,'https://agency.gov.cn/policy/excluded.html',status='excluded',title='合成非政策')
    reference=tmp_path/'reference.jsonl';reference.write_text(json.dumps({'page_url':'https://agency.gov.cn/policy/missed.html','title':'独立列表中的文件'}))
    out=tmp_path/'sample';report=prepare_sample(pipe.db,out,20,'frozen',str(reference))
    assert report['stage_counts']=={'excluded':1,'not_discovered':1}
    score=score_review(out/'sample.jsonl');assert score['human_reviewed']==0 and score['list_recall'] is None
    labels=[json.loads(s) for s in (out/'sample.jsonl').read_text().splitlines()]
    miss=next(g for g in labels if g['pipeline_stage']=='not_discovered');miss.update(relevant=True,categories=['guide'],label_status='human_reviewed',reviewer='测试',reviewed_at='2026-09-10')
    (out/'sample.jsonl').write_text(''.join(json.dumps(g)+'\n' for g in labels))
    assert score_review(out/'sample.jsonl')['relevant_sample_outcomes']['list_not_discovered']==1
    (out/'corpus.jsonl').write_text('tampered')
    with pytest.raises(ValueError,match='快照'):score_review(out/'sample.jsonl')


def test_quality_page_renders_format_counts_not_counter_methods(pipe,monkeypatch):
    """材料质量页按格式统计必须显示数字：Counter 的 total()/属性同名会让模板输出方法对象。"""
    from policy_collector.webapp import create_app
    monkeypatch.setattr(pipe.collector,'fetch',lambda u:FetchResult(content=word(),final_url=u))
    pipe.ingest_url(pipe.cfg.sources['test'],'https://agency.gov.cn/policy/1.html',raw=html(),prefer='rule')
    report=attachment_report(pipe.db)
    assert isinstance(report['by_format']['docx'],dict)
    assert report['by_format']['docx']['total']==1
    body=create_app(pipe.cfg).test_client().get('/quality').get_data(as_text=True)
    assert 'bound method' not in body
    assert 'DOCX' in body and '<td>1</td>' in body


def test_batch_blockers_prevent_model_calls(pipe,tmp_path):
    pipe.sync_sources();sid=pipe.db.get_source('test')['id']
    pipe.db.add_fetch(sid,'https://agency.gov.cn/policy/x.html',title='未下载')
    out=tmp_path/'sample';prepare_sample(pipe.db,out,seed='batch')
    g=json.loads((out/'sample.jsonl').read_text())
    fake=SimpleNamespace(cfg=pipe.cfg,llm=SimpleNamespace(available=False,SYSTEM_PROMPT_TMPL='test'),classify=lambda *a:pytest.fail('不应调用模型'))
    report=run_batch(out,tmp_path/'result',fake,split=g['split'],dry_run=True)
    assert report['status']=='blocked' and report['blockers']


def test_reparse_with_cosmetic_html_change(pipe,monkeypatch):
    original=pipe.parser.parse;raw=word()
    monkeypatch.setattr(pipe.collector,'fetch',lambda u:FetchResult(content=raw,final_url=u))
    monkeypatch.setattr(pipe.parser,'parse',lambda data,fmt,**kw:Document(parse_error='旧解析器') if fmt=='docx' else original(data,fmt,**kw))
    src=pipe.cfg.sources['test'];url='https://agency.gov.cn/policy/1.html'
    pipe.ingest_url(src,url,raw=html(),prefer='rule');before=pipe.db.query_policies()[0]
    monkeypatch.setattr(pipe.parser,'parse',original)
    stats=pipe.ingest_url(src,url,raw=html()+b'<!-- new page counter -->',prefer='rule')
    assert stats.reparsed==1 and len(pipe.db.versions(before['policy_key']))==1
    assert pipe.db.list_attachments(before['id'])[0]['parse_status']=='ok'


def test_deleted_sample_cannot_shrink_denominator(pipe,tmp_path):
    pipe.sync_sources();sid=pipe.db.get_source('test')['id']
    pipe.db.add_fetch(sid,'https://agency.gov.cn/policy/1.html')
    out=tmp_path/'sample';prepare_sample(pipe.db,out)
    (out/'sample.jsonl').write_text('')
    with pytest.raises(ValueError,match='删除'):score_review(out/'sample.jsonl')


def test_batch_validated_labels_and_reported_cost(pipe,tmp_path):
    src=pipe.cfg.sources['test'];url='https://agency.gov.cn/policy/funding.html'
    pipe.ingest_url(src,url,raw=html().replace(b'<a href="a.docx">',b'<span>').replace(b'</a>',b'</span>'),prefer='rule')
    out=tmp_path/'sample';prepare_sample(pipe.db,out)
    g=json.loads((out/'sample.jsonl').read_text());g.update(label_status='human_reviewed',relevant=True,categories=['guarantee'],reviewer='合成测试',reviewed_at='2026-09-10')
    (out/'sample.jsonl').write_text(json.dumps(g)+'\n')
    fake=SimpleNamespace(cfg=pipe.cfg,llm=SimpleNamespace(available=True,SYSTEM_PROMPT_TMPL='test'),
        classify=lambda *a:Classification(is_investment_policy='yes',category='guarantee',method='llm',need_review=False,
            input_tokens=100,output_tokens=20,usage_reported=True))
    report=run_batch(out,tmp_path/'result',fake,split=g['split'],input_rate=1,output_rate=2)
    assert report['status']=='human_reviewed' and report['completed']==1 and report['pending_review_ratio']==0
    assert report['estimated_cost_cny']==pytest.approx(.00014)
    assert (tmp_path/'result/predictions.jsonl').is_file()
